import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Styles visuels
NAVY_HEADER = "1B365D"
HEADER_TEXT = "FFFFFF"
ZEBRA_FILL = "F0F4F8"
SUMMARY_FILL = "E6ECF2"
BORDER_COLOR = "D9D9D9"

font_title = Font(name="Calibri", size=15, bold=True, color=NAVY_HEADER)
font_subtitle = Font(name="Calibri", size=10, italic=True, color="555555")
font_header = Font(name="Calibri", size=11, bold=True, color=HEADER_TEXT)
font_bold = Font(name="Calibri", size=11, bold=True)
font_regular = Font(name="Calibri", size=11)

fill_header = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
fill_zebra = PatternFill(start_color=ZEBRA_FILL, end_color=ZEBRA_FILL, fill_type="solid")
fill_summary = PatternFill(start_color=SUMMARY_FILL, end_color=SUMMARY_FILL, fill_type="solid")

thin_border = Side(border_style="thin", color=BORDER_COLOR)
thick_top = Side(border_style="thin", color=NAVY_HEADER)
double_bottom = Side(border_style="double", color=NAVY_HEADER)

cell_border = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)
summary_border = Border(top=thick_top, bottom=double_bottom, left=thin_border, right=thin_border)

align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")


def generer_excel(data: dict, output_path: str) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Relevé Extrait"
    ws.views.sheetView[0].showGridLines = True

    # # Bloc Entête
    # ws['B2'] = f"RELEVÉ BANCAIRE - {data.get('banque', 'N/C').upper()}"
    # ws['B2'].font = font_title
    # ws['B3'] = f"Titulaire : {data.get('titulaire', 'N/C')} | Période : {data.get('periode', 'N/C')}"
    # ws['B3'].font = font_subtitle

    # En-têtes du tableau
    headers = ["Date de la Valeur", "Type de paiement", "Libellé / Description", "Débit (sortie)", "Crédit (entrée)"]
    start_row = 5
    start_col = 2

    for col_idx, header in enumerate(headers, start=start_col):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = cell_border

    ws.row_dimensions[start_row].height = 25

    # Injection des données
    current_row = start_row + 1
    transactions = data.get("transactions", [])

    for idx, tx in enumerate(transactions):
        c_date = ws.cell(row=current_row, column=2, value=tx.get("date_valeur"))
        c_type = ws.cell(row=current_row, column=3, value=tx.get("type_paiement"))
        c_lib = ws.cell(row=current_row, column=4, value=tx.get("libelle"))
        c_deb = ws.cell(row=current_row, column=5, value=tx.get("debit"))
        c_cre = ws.cell(row=current_row, column=6, value=tx.get("credit"))

        c_date.alignment = align_center
        c_type.alignment = align_left
        c_lib.alignment = align_left
        c_deb.alignment = align_right
        c_cre.alignment = align_right

        c_deb.number_format = '#,##0.00'
        c_cre.number_format = '#,##0.00'

        is_even = (idx % 2 == 1)
        for c in [c_date, c_type, c_lib, c_deb, c_cre]:
            c.font = font_regular
            c.border = cell_border
            if is_even:
                c.fill = fill_zebra

        ws.row_dimensions[current_row].height = 20
        current_row += 1

    # Ligne de Total avec formules Excel
    ws.cell(row=current_row, column=2, value="TOTAL")
    ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=4)
    ws.cell(row=current_row, column=2).alignment = Alignment(horizontal="right", vertical="center")

    total_deb = ws.cell(row=current_row, column=5, value=f"=SUM(E{start_row+1}:E{current_row-1})")
    total_cre = ws.cell(row=current_row, column=6, value=f"=SUM(F{start_row+1}:F{current_row-1})")

    total_deb.number_format = '#,##0.00'
    total_cre.number_format = '#,##0.00'

    for col_idx in range(2, 7):
        c = ws.cell(row=current_row, column=col_idx)
        c.font = font_bold
        c.fill = fill_summary
        c.border = summary_border

    ws.row_dimensions[current_row].height = 24

    # Largeurs de colonnes
    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = 16
    ws.column_dimensions['C'].width = 24
    ws.column_dimensions['D'].width = 55
    ws.column_dimensions['E'].width = 20
    ws.column_dimensions['F'].width = 20

    ws.freeze_panes = 'B6'

    wb.save(output_path)
    return output_path